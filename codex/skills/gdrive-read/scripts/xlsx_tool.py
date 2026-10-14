"""List the sheets of an .xlsx file, or write one sheet (optionally an A1 range) as CSV.

Called by gdrive-read.sh through `uv run --with openpyxl --with defusedxml python3 -I`.
The input is a downloaded, untrusted file: it is only parsed, never executed.
"""

import argparse
import csv
import sys

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


def cmd_list(wb):
    out = csv.writer(sys.stdout, delimiter="\t", lineterminator="\n")
    out.writerow(["index", "title", "rows", "cols"])
    for i, ws in enumerate(wb.worksheets, 1):
        # read_only sheets are unsized when the file has no <dimension>; scan the rows instead.
        if not (ws.max_row and ws.max_column):
            try:
                ws.calculate_dimension(force=True)
            except NameError:  # openpyxl raises UnboundLocalError on a sheet with no rows
                pass
        out.writerow([i, ws.title, ws.max_row or 0, ws.max_column or 0])


def cmd_csv(wb, dest, sheet, a1):
    if sheet:
        if sheet not in wb.sheetnames:
            sys.exit(f"xlsx_tool: no sheet named {sheet!r}")
        ws = wb[sheet]
    else:
        ws = wb.worksheets[0]
    bounds = {}
    if a1:
        min_col, min_row, max_col, max_row = range_boundaries(a1)
        bounds = {"min_col": min_col, "min_row": min_row, "max_col": max_col, "max_row": max_row}
    with open(dest, "w", newline="", encoding="utf-8") as f:
        out = csv.writer(f)
        for row in ws.iter_rows(values_only=True, **bounds):
            out.writerow(["" if v is None else v for v in row])


def main():
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="cmd", required=True)
    p_list = sub.add_parser("list")
    p_list.add_argument("xlsx")
    p_csv = sub.add_parser("csv")
    p_csv.add_argument("xlsx")
    p_csv.add_argument("dest")
    p_csv.add_argument("--sheet")
    p_csv.add_argument("--range")
    args = parser.parse_args()

    wb = load_workbook(args.xlsx, read_only=True, data_only=True)
    try:
        if args.cmd == "list":
            cmd_list(wb)
        else:
            cmd_csv(wb, args.dest, args.sheet, args.range)
    finally:
        wb.close()


if __name__ == "__main__":
    main()
